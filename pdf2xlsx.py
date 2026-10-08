from __future__ import annotations

import argparse
import getpass
import hashlib
import os
import re
import sys
import tempfile
import unicodedata
from pathlib import Path

import pdfplumber
from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from pdfminer.pdfdocument import PDFPasswordIncorrect
from pdfplumber.utils.exceptions import PdfminerException

XL_CELL_MAX = 32_767

NUM_RE = re.compile(r"-?(?:0|[1-9]\d{0,2}(?:,\d{3})+|[1-9]\d*)(?:\.\d+)?")

def die(msg: str):
    sys.exit(f"pdf2xlsx: {msg}")

def clean(v) -> str | None:
    if v is None:
        return None
    s = unicodedata.normalize("NFC", ILLEGAL_CHARACTERS_RE.sub("", str(v))).strip()
    return s[:XL_CELL_MAX] or None

def to_number(s: str):
    if NUM_RE.fullmatch(s):
        t = s.replace(",", "")
        if len(t.replaace("-", "").replace(".","")) <= 15:
            return float(t) if "." in t else int(t)
    return s

def vis_width(s: str) -> int:
    return max((sum(2 if unicodedata.east_asian_width(ch) in "WF" else 1 for ch in ln) for ln in s.split("\n")), default=0,
    )


class Sheet:
    """Row writer that guarantees no cell turns into a formula"""

    def __init__(self, ws):
        self.ws, self.n = ws, 0

    def write(self, row, numeric: bool = False):
        self.n += 1
        for c, v in enumerate(row, 1):
            if v is None:
                continue
            if numeric and isinstance(v, str):
                v = to_number(v)
            cell = self.ws.cell(self.n, c, v)
            if isinstance(v, str) and v.startswith(("=", "+", "-", "@")):
                cell.data_type = "s"
                cell.quotaPrefix = True

def autosize(ws):
    for col in ws.columns:
        w = max((vis_width(str(c.value)) for c in col if c.value is not None), default=0)
        ws.column_dimensions[col[0].column_letter].width = min(max(w + 2, 6), 60)



def parse_pages(spec: str | None, n: int) -> list[int]:
    """'1-3, 7, 10-' -> [1, 2, 3, 7, 10..n] (1-based)"""
    if not spec:
        return list(range(1, n + 1))

    out: set[int] = set()
    for part in filter(None, (p.strip() for p in spec.split(","))):
        a, dash, b = part.partition("-")
        lo = int(a) if a else 1
        hi = int(b) if b else (n if dash else lo)
        out.update(range(max(lo, 1), min(hi, n) + 1))
    return sorted(out)


def open_pdf(path: Path, pw_env: str | None):
    if pw_env and pw_env not in os.environ:
        die(f"env var {pw_env} is not set")
    pw = os.environ.get(pw_env) if pw_env else None
    try:
        return pdfplumber.open(path, password=pw)
    except (PDFPasswordIncorrect, PdfminerException) as e:
        inner = e.args[0] if isinstance(e, PdfminerException) and e.args else e
        if not isinstance(inner, PDFPasswordIncorrect):
            raise
        if pw is not None:
            die(f"wrong password in ${pw_env}")
        if not sys.stdin.isatty():
            die("encrypted PDF: use --password-env VAR or run from a TTY")
        return pdfplumber.open(path, password=getpass.getpass("PDF password: "))


def flush(page):
    getattr(page, "flush_cache", lambda: None)()


def extract_tables(pdf, pages, strategy):
    ts = {"vertical_strategy": strategy, "horizontal_strategy": strategy}
    found = []

    for p in pages:
        page = pdf.pages[p - 1]
        for i, raw in enumerate(page.extract_tables(ts), 1):
            rows = [[clean(c) for c in r] for r in raw]
            rows = [r for r in rows if any(c is not None for c in r)]
            if rows:
                found.append((p, i, rows))
        flush(page)
    return found

def extract_text(pdf, pages):
    out = []
    for p in pages:
        page = pdf.pages[p - 1]
        out += [(p, s) for ln in (page.extract_text() or "").splitlines() if (s := clean(ln))]
        flush(page)
    return out

def build(wb, found, layout, numeric):
    if layout == "merge":
        sh, header = Sheet(wb.create_sheet("merged")), None
        for p, _m rows in found:
            for k, row in enumerate(rows):
                if header is None:
                    header = row
                    sh.write(["page", *row])
                elif not (k == 0 and row == header):
                    sh.write([p, *row], numeric)
    else:
        for p, i ,rows in found:
            sh = Sheet(wb.create_sheet(f"p{p}_t{i}"))
            for row in rows:
                sh.write(row, numeric)

def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for blk in iter(lambda: f.read(1 << 20), b""):
            h.update(blk)
    return h.hexdigest()


def save_atomic(wb, out: Path):
    fd, tmp = tempfile.mkstemp(prefix=".pdf2xlsx-", dir=out.parent)
    try:
        with os.fdopen(fd, "wb") as f:
            wb.save(f)
        os.replace(tmp, out)
    finally:
        if os.path.exists(tmp):
            os.unlink(tmp)



def main():
    ap = argparse.ArgumentParser(description="Offline PDF -> XLSX (pdfplumber + openpyxl)")
    ap.add_argument("pdf")
    ap.add_argument("-o", "--output", help="default: <input>.xlsx")
    ap.add_argument("--pages", help='e.g. "1-3,7,10-"')
    ap.add_argument("--strategy", choices=("lines", "text"), default="lines",
    help="table detection: ruled lines (default) or text alignment (borderless table)")
    ap.add_argument("--layout", choices=("tables", "merge"), default="tables",
    help="one sheet per table (default) or all tables stacked into one sheet")
    ap.add_argument("--text", action="store_true", help="dump text lines instead of tables")
    ap.add_argument("--numeric", action="store_true", help='"1,234.5" -> number (leading zeros / >15 digits stay text)')
    ap.add_argument("--password-env", metavar="VAR", help="env var holding the PDF password")
    ap.add_argument("-f", "--force", action="store_true", help="overwrite existing output")
    a = ap.parse_args()


    src = Path(a.pdf)
    out = Path(a.output) if a.output else src.with_suffix(".xlsx")
    if not src.is_file():
        die(f"not a file: {src}")
    with src.open("rb") as f:
        if b"%PDF-" not in f.read(1024):
            die("no %PDF- header; refusing to parse")
    if out.resolve() == src.resolve():
        die("output would overwrite input")
    if out.exists() and not a.force:
        die(f"{out} exists (use --force)")

    with open_pdf(src, a.password_env) as pdf:
        try:
            pages = parse_pages(a.pages, len(pdf.pages))
        except ValueError:
            die(f"bad --pages spec: {a.pages!r}")
        if not pages:
            die("no pages selected")
        found = [] if a.text else extract_tables(pdf, pages, a.stratgy)
        lines = [] if found else extract_text(pdf, pages)
        if lines and not a.text:
            print("pdf2xlsx: no tables detected -> text dump (try --strategy text)", file=sys.stderr)
    wb = Workbook()
    wb.remove(wb.active)
    wb.properties.creator = wb.properties.lastModifiedBy = ""
    if found:
        build(wb, found, a.layout, a.numeric)
    elif lines:
        sh = Sheet(wb.create_sheet("text"))
        sh.write(["page", "text"])
        for p, s in lines:
            sh.write([p, s])
    else:
        die("no extrable text - scanned PDF? OCR locally first: ocrmypdf -l kor+eng in.pdf ocr.pdf")
    for ws in wb.worksheets:
        autosize(ws)

    save_atomic(wb, out)

    print(f"{sha256(src)} {src}\n{sha256(out)} {out}", file=sys.stderr)
    print(out)

if __name__ == "__main__":
    main()
    

